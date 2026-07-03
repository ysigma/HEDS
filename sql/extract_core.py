"""
Enriched template-metadata extraction (v3).

The functions below drop verbatim into the Snowflake stored procedure handler;
only the Snowflake I/O (download, DELETE / INSERT) lives outside
`extract_workbook_metadata`. The procedure body is dollar-quoted, so
apostrophes here are safe.

What it captures, across the whole workbook
-------------------------------------------
* Excel Tables (structured), including multi-header-row tables and single-row
  "wide record" tables (also decomposed into per-field scalars).
* Defined names (workbook-global and sheet-local): scalars, 1-D lists and 2-D
  ranges; broken (#REF!), built-in (_xlnm.*) and full-column names are skipped.
* Key/value fact blocks (horizontal and vertical), fund identifiers, fundamentals.
* Detected breakdown tables that are NOT Excel Tables (connected-component scan,
  with spacer-column de-fragmentation and title-row peeling).
* Footnotes / labeled long text, and "Label: value" single cells.

For every element it records sheet, cell/table/header/data refs, orientation,
row/column counts, header names, per-column and scalar value KINDS (number-format
aware: PERCENT / CURRENCY / INTEGER / DATE / TEXT ...), the current cached value
and the underlying formula.

Downstream compatibility
------------------------
The Sigma plugin matches metadata COLUMNS to roles by (normalized) name -- only
SHEET_NAME / ELEMENT_TYPE / ELEMENT_NAME / CELL_REF / TABLE_REF / COLUMNS_JSON
are consumed; a type containing "TABLE" serializes as an array, "NAMED" (or no
columns) as a scalar. Every other column is ignored -- but a name normalizing to
contain "column" hijacks the columns role and one containing "type" hijacks the
elementType role, so the extra columns are deliberately named NUM_COLS /
NUM_ROWS / COL_KINDS_JSON / VALUE_KIND / LABEL_REF / HEADER_REF / DATA_REF etc.
"""

import re
import json
import datetime
from openpyxl.utils import range_boundaries, get_column_letter

try:
    from openpyxl.worksheet.worksheet import Worksheet
except Exception:                       # pragma: no cover
    Worksheet = None

try:
    from openpyxl.styles.numbers import is_date_format
except Exception:                       # pragma: no cover
    def is_date_format(fmt):
        f = (fmt or "").lower()
        return any(t in f for t in ("yy", "mm", "dd", "hh", "ss"))


MAX_TEXT = 2000            # cap stored raw text
MERGE_MAX_GAP = 1          # spacer columns bridged when merging table fragments
MIN_TABLE_ROWS_FOR_MERGE = 3
SENTENCE_MIN_LEN = 45      # text at/above this is treated as prose
MAX_HEADER_ROWS = 2        # leading label rows composed into one header


# --------------------------------------------------------------------------
# Value helpers
# --------------------------------------------------------------------------

def is_formula(v):
    return isinstance(v, str) and v.startswith("=")


def is_blank(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


def is_number(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _numeric_text(s):
    num = s.strip().strip("()")
    for ch in ("%", ",", "$", "€", "£", "¥", " "):
        num = num.replace(ch, "")
    try:
        float(num)
        return True
    except ValueError:
        return False


def is_label_text(v):
    """Text that reads like a field label / header (non-empty, non-numeric).
    Text-formatted numbers (percent, currency, thousands, accounting negatives)
    are values, not labels."""
    if not isinstance(v, str):
        return False
    s = v.strip()
    if s == "" or s.startswith("="):
        return False
    return not _numeric_text(s)


def is_sentence_like(v):
    """Prose (footnote / disclaimer / note), not a short label."""
    if not isinstance(v, str):
        return False
    s = v.strip()
    if len(s) >= SENTENCE_MIN_LEN:
        return True
    return (" " in s) and (s.endswith(".") or s.count(" ") >= 6)


def value_kind(v):
    if is_blank(v):
        return "EMPTY"
    if isinstance(v, bool):
        return "BOOLEAN"
    if isinstance(v, (datetime.datetime, datetime.date)):
        return "DATE"
    if is_number(v):
        return "NUMBER"
    return "TEXT"


def format_kind(number_format, cached):
    """Refine a cached value's kind using the cell's number format so percent /
    currency / integer are distinguished from a bare NUMBER."""
    base = value_kind(cached)
    fmt = number_format or "General"
    if base == "NUMBER":
        if "%" in fmt:
            return "PERCENT"
        if is_date_format(fmt):
            return "DATE"
        if any(sym in fmt for sym in ("$", "€", "£", "¥", "[$")):
            return "CURRENCY"
        core = fmt.split(";")[0]
        if "0" in core and "." not in core and "E" not in core.upper():
            return "INTEGER"
        return "NUMBER"
    if base == "TEXT" and is_date_format(fmt):
        return "DATE"
    return base


def to_text(v):
    if v is None:
        return None
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.isoformat()
    s = v if isinstance(v, str) else str(v)
    return s[:MAX_TEXT]


def clean_name(v):
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v)).strip()


def dedupe_headers(headers):
    """Ensure column names are unique so a downstream row object can't lose a
    duplicated column (e.g. two 'Total' columns)."""
    seen, out = {}, []
    for i, h in enumerate(headers):
        h = h if h else f"Column{i+1}"
        if h in seen:
            seen[h] += 1
            out.append(f"{h} ({seen[h]})")
        else:
            seen[h] = 1
            out.append(h)
    return out


# --------------------------------------------------------------------------
# Range helpers (robust to malformed refs)
# --------------------------------------------------------------------------

FULL_RANGE_RE = re.compile(r"^\$?[A-Za-z]+:\$?[A-Za-z]+$|^\$?\d+:\$?\d+$")


def safe_bounds(ref):
    """range_boundaries() but returns None for #REF!, full-column/row, external
    and other refs that can't be turned into a finite rectangle."""
    if not ref or "#REF!" in ref or "!" in ref or "[" in ref:
        return None
    if FULL_RANGE_RE.match(ref.strip()):
        return None
    try:
        c1, r1, c2, r2 = range_boundaries(ref.replace("$", ""))
    except Exception:
        return None
    if None in (c1, r1, c2, r2):
        return None
    return c1, r1, c2, r2


def cell_ref(row, col):
    return f"{get_column_letter(col)}{row}"


def rect_ref(r1, c1, r2, c2):
    return f"{cell_ref(r1, c1)}:{cell_ref(r2, c2)}"


def norm_ref(ref):
    return ref.replace("$", "") if isinstance(ref, str) else ref


def kind_at(ws, ws_vals, r, c):
    return format_kind(ws.cell(row=r, column=c).number_format,
                       ws_vals.cell(row=r, column=c).value)


def infer_column_kinds(ws, ws_vals, first_body_row, last_row, c1, c2):
    kinds = []
    for c in range(c1, c2 + 1):
        seen = set()
        for r in range(first_body_row, last_row + 1):
            v = ws_vals.cell(row=r, column=c).value
            if not is_blank(v):
                seen.add(kind_at(ws, ws_vals, r, c))
        if not seen:
            kinds.append("EMPTY")
        elif len(seen) == 1:
            kinds.append(next(iter(seen)))
        else:
            kinds.append("MIXED")
    return kinds


# --------------------------------------------------------------------------
# Emitter
# --------------------------------------------------------------------------

class Emitter:
    def __init__(self):
        self.rows = []
        self._seq = {}
        self._seen = set()
        self.stats = {
            "skipped_single_cells": 0,
            "skipped_examples": [],
            "skipped_sheets": [],
            "sheet_errors": [],
            "empty_yield_sheets": [],
        }

    def emit(self, sheet, **kw):
        name = kw.get("ELEMENT_NAME") or "?"
        dedup = (sheet, name, kw.get("CELL_REF"), kw.get("ELEMENT_TYPE"))
        if dedup in self._seen:
            return
        self._seen.add(dedup)
        seq = self._seq.get(sheet, 0)
        self._seq[sheet] = seq + 1
        row = {k: None for k in (
            "ELEMENT_KEY", "ELEMENT_TYPE", "ELEMENT_NAME", "SHEET_NAME",
            "CELL_REF", "LABEL_REF", "TABLE_REF", "HEADER_REF", "DATA_REF",
            "ORIENTATION", "NUM_ROWS", "NUM_COLS", "COLUMNS_JSON",
            "COL_KINDS_JSON", "VALUE_KIND", "RAW_VALUE", "FORMULA",
            "DETECTION_METHOD", "SEQ")}
        row["SHEET_NAME"] = sheet
        row["SEQ"] = seq
        row.update(kw)
        row["ELEMENT_KEY"] = f"{sheet}!{name}#{seq}"
        self.rows.append(row)

    def note_skipped(self, sheet, coord, value):
        self.stats["skipped_single_cells"] += 1
        if len(self.stats["skipped_examples"]) < 20:
            self.stats["skipped_examples"].append(f"{sheet}!{coord}={to_text(value)!r}"[:90])


# --------------------------------------------------------------------------
# Excel Tables
# --------------------------------------------------------------------------

def _compose_headers(ws, r1, c1, c2, depth):
    cols = []
    for c in range(c1, c2 + 1):
        parts = []
        for r in range(r1, r1 + depth):
            v = ws.cell(row=r, column=c).value
            if not is_blank(v):
                parts.append(clean_name(v))
        cols.append(" | ".join(parts))
    return cols


def extract_tables(emitter, ws, ws_vals):
    for tn in list(ws.tables):
        table = ws.tables[tn]
        b = safe_bounds(table.ref)
        if b is None:
            continue
        c1, r1, c2, r2 = b
        header_rows = table.headerRowCount if table.headerRowCount is not None else 1
        header_rows = max(0, header_rows)
        totals_rows = table.totalsRowCount or 0
        if header_rows > 0:
            columns = _compose_headers(ws, r1, c1, c2, header_rows)
        else:
            columns = [""] * (c2 - c1 + 1)
        columns = dedupe_headers(
            [clean_name(x) if not is_blank(x) else f"Column{i+1}"
             for i, x in enumerate(columns)])
        body_top = r1 + header_rows
        body_bot = r2 - totals_rows
        data_rows = max(0, body_bot - body_top + 1)
        emitter.emit(
            ws.title,
            ELEMENT_TYPE="TABLE",
            ELEMENT_NAME=clean_name(table.displayName or tn),
            CELL_REF=cell_ref(body_top, c1),
            TABLE_REF=table.ref,
            HEADER_REF=(rect_ref(r1, c1, r1 + header_rows - 1, c2) if header_rows else None),
            DATA_REF=(rect_ref(body_top, c1, body_bot, c2) if data_rows > 0 else None),
            ORIENTATION="vertical",
            NUM_ROWS=data_rows, NUM_COLS=(c2 - c1 + 1),
            COLUMNS_JSON=json.dumps(columns),
            COL_KINDS_JSON=json.dumps(
                infer_column_kinds(ws, ws_vals, body_top, body_bot, c1, c2)),
            DETECTION_METHOD="EXCEL_TABLE",
        )
        # A single-data-row, multi-field table is a "wide record" (e.g. the fund
        # data block). Also expose each field as a scalar so it maps uniformly
        # with sheets that lay the same fields out as loose key/value cells.
        if data_rows == 1 and header_rows == 1 and (c2 - c1 + 1) >= 3:
            for i, c in enumerate(range(c1, c2 + 1)):
                label = ws.cell(row=r1, column=c).value
                if not is_label_text(label):
                    continue
                emitter.emit(
                    ws.title,
                    ELEMENT_TYPE="KEY_VALUE",
                    ELEMENT_NAME=clean_name(label),
                    CELL_REF=cell_ref(body_top, c), LABEL_REF=cell_ref(r1, c),
                    ORIENTATION="horizontal", NUM_ROWS=1, NUM_COLS=1,
                    VALUE_KIND=kind_at(ws, ws_vals, body_top, c),
                    RAW_VALUE=to_text(ws_vals.cell(row=body_top, column=c).value),
                    DETECTION_METHOD="WIDE_RECORD",
                )


# --------------------------------------------------------------------------
# Defined names
# --------------------------------------------------------------------------

def _skip_defined_name(name, refers_to):
    if refers_to is None or "#REF!" in refers_to:
        return True
    low = name.lower()
    if low.startswith("_xlnm") or low.startswith("_xlfn") or name.startswith("_"):
        return True
    if "print_area" in low or "print_titles" in low:
        return True
    return False


def _iter_defined_names(wb, scope_ws):
    container = wb.defined_names if scope_ws is None else scope_ws.defined_names
    for name in list(container):
        dn = container[name]
        if _skip_defined_name(name, dn.attr_text or ""):
            continue
        try:
            dests = list(dn.destinations)
        except Exception:
            continue
        for sheet_name, coord in dests:
            if not coord or sheet_name not in wb.sheetnames:
                continue
            b = safe_bounds(coord)
            if b is None:
                continue
            yield name, sheet_name, coord, b


def extract_defined_names(emitter, wb, wb_vals, scope_ws=None):
    for name, sheet_name, coord, (c1, r1, c2, r2) in _iter_defined_names(wb, scope_ws):
        ws, ws_vals = wb[sheet_name], wb_vals[sheet_name]
        nrows, ncols = r2 - r1 + 1, c2 - c1 + 1
        if nrows == 1 and ncols == 1:
            raw = ws.cell(row=r1, column=c1).value
            emitter.emit(
                sheet_name,
                ELEMENT_TYPE="NAMED_RANGE",
                ELEMENT_NAME=clean_name(name),
                CELL_REF=norm_ref(coord),
                ORIENTATION="scalar", NUM_ROWS=1, NUM_COLS=1,
                VALUE_KIND=kind_at(ws, ws_vals, r1, c1),
                RAW_VALUE=to_text(ws_vals.cell(row=r1, column=c1).value),
                FORMULA=(raw if is_formula(raw) else None),
                DETECTION_METHOD="DEFINED_NAME",
            )
        elif ncols > 1 and nrows > 1:
            headers = dedupe_headers(_compose_headers(ws, r1, c1, c2, 1))
            emitter.emit(
                sheet_name,
                ELEMENT_TYPE="NAMED_TABLE",
                ELEMENT_NAME=clean_name(name),
                CELL_REF=cell_ref(r1 + 1, c1),
                TABLE_REF=rect_ref(r1, c1, r2, c2),
                HEADER_REF=rect_ref(r1, c1, r1, c2),
                DATA_REF=rect_ref(r1 + 1, c1, r2, c2),
                ORIENTATION="vertical", NUM_ROWS=nrows - 1, NUM_COLS=ncols,
                COLUMNS_JSON=json.dumps(headers),
                COL_KINDS_JSON=json.dumps(
                    infer_column_kinds(ws, ws_vals, r1 + 1, r2, c1, c2)),
                DETECTION_METHOD="DEFINED_NAME",
            )
        else:
            # 1-D range -> a list. Type contains TABLE so downstream serializes
            # an array (every row), not just the first cell.
            emitter.emit(
                sheet_name,
                ELEMENT_TYPE="NAMED_TABLE",
                ELEMENT_NAME=clean_name(name),
                CELL_REF=cell_ref(r1, c1),
                TABLE_REF=rect_ref(r1, c1, r2, c2),
                DATA_REF=rect_ref(r1, c1, r2, c2),
                ORIENTATION=("row" if nrows == 1 else "column"),
                NUM_ROWS=nrows, NUM_COLS=ncols,
                COLUMNS_JSON=json.dumps([clean_name(name)]),
                DETECTION_METHOD="DEFINED_NAME",
            )


# --------------------------------------------------------------------------
# Region detection
# --------------------------------------------------------------------------

def collect_claimed(wb, ws):
    """Cells that belong to an Excel Table or a MULTI-cell defined name on this
    sheet -- the region scan skips them. Single-cell named ranges are left
    unclaimed on purpose: they usually anchor a label whose value sits in an
    adjacent cell, which the key/value scan must still pair up."""
    claimed = set()
    for tn in list(ws.tables):
        b = safe_bounds(ws.tables[tn].ref)
        if b is None:
            continue
        c1, r1, c2, r2 = b
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                claimed.add((r, c))
    for scope in (None, ws):
        for _name, sheet_name, _coord, (c1, r1, c2, r2) in _iter_defined_names(wb, scope):
            if sheet_name != ws.title:
                continue
            if (r2 - r1 + 1) * (c2 - c1 + 1) > 1:
                for r in range(r1, r2 + 1):
                    for c in range(c1, c2 + 1):
                        claimed.add((r, c))
    return claimed


def build_merged_map(ws):
    covered = {}
    for m in ws.merged_cells.ranges:
        r1, c1 = m.min_row, m.min_col
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                if (r, c) != (r1, c1):
                    covered[(r, c)] = (r1, c1)
    return covered


def connected_components(cells):
    cells = set(cells)
    seen = set()
    comps = []
    for start in cells:
        if start in seen:
            continue
        stack, comp = [start], []
        seen.add(start)
        while stack:
            r, c = stack.pop()
            comp.append((r, c))
            for dr, dc in ((1, 0), (-1, 0), (0, 1), (0, -1)):
                nb = (r + dr, c + dc)
                if nb in cells and nb not in seen:
                    seen.add(nb)
                    stack.append(nb)
        comps.append(set(comp))
    return comps


def _bbox(cells):
    rs = [r for r, _ in cells]
    cs = [c for _, c in cells]
    return min(rs), min(cs), max(rs), max(cs)


def merge_table_fragments(comps):
    """Merge tall components separated by <= MERGE_MAX_GAP spacer columns that
    share most of their row span (a table split by a blank column). Short strips
    / single cells never merge, so annotations stay separate."""
    boxes = [{"cells": set(c), "box": _bbox(c)} for c in comps]

    def tall(o):
        r1, _c1, r2, _c2 = o["box"]
        return (r2 - r1 + 1) >= MIN_TABLE_ROWS_FOR_MERGE

    changed = True
    while changed:
        changed = False
        for i in range(len(boxes)):
            if boxes[i] is None:
                continue
            for j in range(i + 1, len(boxes)):
                if boxes[j] is None or not (tall(boxes[i]) and tall(boxes[j])):
                    continue
                a, b = boxes[i], boxes[j]
                ar1, _, ar2, ac2 = a["box"]
                br1, bc1, br2, _ = b["box"]
                overlap = min(ar2, br2) - max(ar1, br1) + 1
                if overlap <= 0 or overlap < 0.5 * min(ar2 - ar1 + 1, br2 - br1 + 1):
                    continue
                left, right = (a, b) if a["box"][3] <= b["box"][1] else (b, a)
                gap = right["box"][1] - left["box"][3] - 1
                if 0 <= gap <= MERGE_MAX_GAP:
                    a["cells"] |= b["cells"]
                    a["box"] = _bbox(a["cells"])
                    boxes[j] = None
                    changed = True
    return [o["cells"] for o in boxes if o is not None]


def _header_depth(ws, r1, c1, c2, last_row):
    """Number of leading all-label rows to treat as a (possibly multi-row)
    header; always >= 1, capped, and never consuming the last data row."""
    depth = 0
    for r in range(r1, min(r1 + MAX_HEADER_ROWS, last_row) + 1):
        nonblank = [ws.cell(row=r, column=c).value for c in range(c1, c2 + 1)]
        nonblank = [v for v in nonblank if not is_blank(v)]
        if nonblank and all(is_label_text(v) for v in nonblank) and r < last_row:
            depth += 1
        else:
            break
    return max(1, depth)


def _emit_data_table(emitter, ws, ws_vals, cells, method):
    r1, c1, r2, c2 = _bbox(cells)
    if r2 <= r1:
        return
    depth = _header_depth(ws, r1, c1, c2, r2)
    body_top = r1 + depth
    headers, kinds, kept = [], [], 0
    for c in range(c1, c2 + 1):
        parts = [clean_name(ws.cell(row=r, column=c).value)
                 for r in range(r1, r1 + depth)
                 if not is_blank(ws.cell(row=r, column=c).value)]
        body_seen = set()
        for r in range(body_top, r2 + 1):
            v = ws_vals.cell(row=r, column=c).value
            if not is_blank(v):
                body_seen.add(kind_at(ws, ws_vals, r, c))
        if not parts and not body_seen:
            continue                              # drop pure spacer column
        headers.append(" | ".join(parts) if parts else f"Column{kept+1}")
        kinds.append("EMPTY" if not body_seen else
                     (next(iter(body_seen)) if len(body_seen) == 1 else "MIXED"))
        kept += 1
    if not headers:
        return
    emitter.emit(
        ws.title,
        ELEMENT_TYPE="DATA_TABLE",
        ELEMENT_NAME=f"{ws.title}!{rect_ref(r1, c1, r2, c2)}",
        CELL_REF=cell_ref(body_top, c1),
        TABLE_REF=rect_ref(r1, c1, r2, c2),
        HEADER_REF=rect_ref(r1, c1, r1 + depth - 1, c2),
        DATA_REF=rect_ref(body_top, c1, r2, c2),
        ORIENTATION="vertical",
        NUM_ROWS=r2 - body_top + 1, NUM_COLS=len(headers),
        COLUMNS_JSON=json.dumps(dedupe_headers(headers)),
        COL_KINDS_JSON=json.dumps(kinds),
        DETECTION_METHOD=method,
    )


def _emit_kv(emitter, ws, ws_vals, name, value_r, value_c, label_r, label_c, orient):
    raw = ws.cell(row=value_r, column=value_c).value
    emitter.emit(
        ws.title,
        ELEMENT_TYPE="KEY_VALUE",
        ELEMENT_NAME=clean_name(name),
        CELL_REF=cell_ref(value_r, value_c),
        LABEL_REF=cell_ref(label_r, label_c),
        ORIENTATION=orient, NUM_ROWS=1, NUM_COLS=1,
        VALUE_KIND=kind_at(ws, ws_vals, value_r, value_c),
        RAW_VALUE=to_text(ws_vals.cell(row=value_r, column=value_c).value),
        FORMULA=(raw if is_formula(raw) else None),
        DETECTION_METHOD="KEY_VALUE_SCAN",
    )


def classify_and_emit(emitter, ws, ws_vals, cells):
    getv = lambda r, c: ws.cell(row=r, column=c).value

    # Peel leading title rows: a single filled cell sitting atop a wider block
    # is a section title, not the table's header row.
    while True:
        r1, c1, r2, c2 = _bbox(cells)
        if r2 <= r1:
            break
        toprow = [(r, c) for (r, c) in cells if r == r1]
        bodycols = {c for (r, c) in cells if r > r1}
        if len(toprow) == 1 and len(bodycols) >= 2:
            tr, tc = toprow[0]
            emitter.note_skipped(ws.title, cell_ref(tr, tc), getv(tr, tc))
            cells = {(r, c) for (r, c) in cells if r != r1}
        else:
            break

    r1, c1, r2, c2 = _bbox(cells)
    nrows, ncols = r2 - r1 + 1, c2 - c1 + 1

    # ---- single cell -> skip, unless it is an inline "Label: value" --
    if nrows == 1 and ncols == 1:
        v = getv(r1, c1)
        if isinstance(v, str) and ":" in v and not is_formula(v) and not is_sentence_like(v):
            left, right = v.split(":", 1)
            left, right = left.strip(), right.strip()
            if right and 0 < len(left) <= 40 and not is_sentence_like(right) \
                    and is_label_text(left) and not _numeric_text(left):
                emitter.emit(
                    ws.title, ELEMENT_TYPE="KEY_VALUE", ELEMENT_NAME=clean_name(left),
                    CELL_REF=cell_ref(r1, c1), LABEL_REF=cell_ref(r1, c1),
                    ORIENTATION="inline", NUM_ROWS=1, NUM_COLS=1,
                    VALUE_KIND="TEXT", RAW_VALUE=to_text(right),
                    DETECTION_METHOD="KEY_VALUE_SCAN")
                return
        emitter.note_skipped(ws.title, cell_ref(r1, c1), v)
        return

    # ---- single-row strip -> pair alternating label/value cells ------
    if nrows == 1:
        c = c1
        while c <= c2:
            if is_label_text(getv(r1, c)) and c + 1 <= c2 and not is_blank(getv(r1, c + 1)) \
                    and not is_label_text(getv(r1, c + 1)):
                _emit_kv(emitter, ws, ws_vals, getv(r1, c), r1, c + 1, r1, c, "horizontal")
                c += 2
            else:
                if not is_blank(getv(r1, c)):
                    emitter.note_skipped(ws.title, cell_ref(r1, c), getv(r1, c))
                c += 1
        return

    # ---- single column -----------------------------------------------
    if ncols == 1:
        nonblank = [getv(r, c1) for r in range(r1, r2 + 1) if not is_blank(getv(r, c1))]
        prose = sum(1 for v in nonblank if is_sentence_like(v))
        if nrows == 2:
            emitter.emit(
                ws.title,
                ELEMENT_TYPE=("LABELED_TEXT" if is_sentence_like(getv(r2, c1)) else "KEY_VALUE"),
                ELEMENT_NAME=clean_name(getv(r1, c1)),
                CELL_REF=cell_ref(r2, c1), LABEL_REF=cell_ref(r1, c1),
                ORIENTATION="vertical", NUM_ROWS=1, NUM_COLS=1,
                VALUE_KIND=kind_at(ws, ws_vals, r2, c1),
                RAW_VALUE=to_text(ws_vals.cell(row=r2, column=c1).value),
                DETECTION_METHOD="REGION_SCAN")
            return
        if prose >= max(1, len(nonblank) - 1):
            for r in range(r1, r2 + 1):
                v = getv(r, c1)
                if is_blank(v):
                    continue
                emitter.emit(
                    ws.title, ELEMENT_TYPE="FOOTNOTE",
                    ELEMENT_NAME=(clean_name(v)[:60] or "Footnote"),
                    CELL_REF=cell_ref(r, c1), ORIENTATION="scalar",
                    NUM_ROWS=1, NUM_COLS=1, VALUE_KIND="TEXT",
                    RAW_VALUE=to_text(ws_vals.cell(row=r, column=c1).value),
                    DETECTION_METHOD="REGION_SCAN")
            return
        _emit_data_table(emitter, ws, ws_vals, cells, "REGION_SCAN")
        return

    # ---- two rows -> horizontal key/value (only if every value has a
    #      header label above it; otherwise it is a real table) --------
    if nrows == 2:
        value_cols = {c for c in range(c1, c2 + 1) if not is_blank(getv(r2, c))}
        label_cols = {c for c in range(c1, c2 + 1) if is_label_text(getv(r1, c))}
        if label_cols and value_cols <= label_cols:
            for c in sorted(label_cols):
                _emit_kv(emitter, ws, ws_vals, getv(r1, c), r2, c, r1, c, "horizontal")
            return
        _emit_data_table(emitter, ws, ws_vals, cells, "REGION_SCAN")
        return

    # ---- two columns: 2-col table (header + numeric body) vs vertical KV
    if ncols == 2:
        header_pair = is_label_text(getv(r1, c1)) and is_label_text(getv(r1, c2))
        body = [ws_vals.cell(row=r, column=c2).value for r in range(r1 + 1, r2 + 1)]
        body = [v for v in body if not is_blank(v)]
        numeric_body = body and sum(
            1 for v in body if value_kind(v) in ("NUMBER", "DATE")) >= 0.6 * len(body)
        if header_pair and numeric_body:
            _emit_data_table(emitter, ws, ws_vals, cells, "REGION_SCAN")
            return
        left_labels = [r for r in range(r1, r2 + 1) if is_label_text(getv(r, c1))]
        col_rows = {r for (r, c) in cells if c == c1}
        if left_labels and len(left_labels) == len(col_rows):
            for r in range(r1, r2 + 1):
                if is_label_text(getv(r, c1)):
                    _emit_kv(emitter, ws, ws_vals, getv(r, c1), r, c2, r, c1, "vertical")
            return
        _emit_data_table(emitter, ws, ws_vals, cells, "REGION_SCAN")
        return

    # ---- otherwise -> detected data table ----------------------------
    _emit_data_table(emitter, ws, ws_vals, cells, "REGION_SCAN")


def extract_regions(emitter, ws, ws_vals, claimed):
    covered = build_merged_map(ws)
    filled = {}
    for row in ws.iter_rows():
        for cell in row:
            rc = (cell.row, cell.column)
            if rc in claimed or rc in covered:
                continue
            if not is_blank(cell.value):
                filled[rc] = cell.value
    if not filled:
        return
    # Bridge connectivity through merged ranges (covered cells) so a merged
    # banner/label can't split one logical block, but only real (filled) cells
    # are ever emitted.
    connect = set(filled)
    for rc, anchor in covered.items():
        if anchor in filled and rc not in claimed:
            connect.add(rc)
    fset = set(filled)
    comps = [c & fset for c in connected_components(connect)]
    comps = [c for c in comps if c]
    for cells in merge_table_fragments(comps):
        classify_and_emit(emitter, ws, ws_vals, cells)


# --------------------------------------------------------------------------
# Orchestration
# --------------------------------------------------------------------------

def _sheet_has_content(ws):
    return ws.max_row and ws.max_column and (ws.max_row > 1 or ws.max_column > 1
                                             or ws.cell(row=1, column=1).value is not None)


def extract_workbook_metadata(wb, wb_vals):
    emitter = Emitter()
    extract_defined_names(emitter, wb, wb_vals, scope_ws=None)   # workbook-global
    for sheet_name in wb.sheetnames:
        try:
            ws = wb[sheet_name]
            if Worksheet is not None and not isinstance(ws, Worksheet):
                emitter.stats["skipped_sheets"].append([sheet_name, "non-worksheet"])
                continue
            if getattr(ws, "sheet_state", "visible") != "visible":
                emitter.stats["skipped_sheets"].append([sheet_name, ws.sheet_state])
                continue
            ws_vals = wb_vals[sheet_name]
            before = len(emitter.rows)
            claimed = collect_claimed(wb, ws)
            extract_tables(emitter, ws, ws_vals)
            extract_defined_names(emitter, wb, wb_vals, scope_ws=ws)  # sheet-local
            extract_regions(emitter, ws, ws_vals, claimed)
            if len(emitter.rows) == before and _sheet_has_content(ws):
                emitter.stats["empty_yield_sheets"].append(sheet_name)
        except Exception as e:                                       # never abort the run
            emitter.stats["sheet_errors"].append(f"{sheet_name}: {e}")

    # Cross-sheet duplicate scalar names collide when a payload keys purely by
    # element name -- surface them (never silent).
    scalar = {}
    for row in emitter.rows:
        if "TABLE" not in (row["ELEMENT_TYPE"] or ""):
            key = clean_name(row["ELEMENT_NAME"]).lower()
            scalar.setdefault(key, set()).add(row["SHEET_NAME"])
    emitter.stats["duplicate_scalar_names"] = sorted(
        n for n, sheets in scalar.items() if len(sheets) > 1)
    return emitter.rows, emitter.stats


if __name__ == "__main__":
    import openpyxl, sys, warnings
    warnings.simplefilter("ignore")
    F = sys.argv[1]
    only = sys.argv[2] if len(sys.argv) > 2 else None
    wb = openpyxl.load_workbook(F, data_only=False)
    wbv = openpyxl.load_workbook(F, data_only=True)
    rows, stats = extract_workbook_metadata(wb, wbv)
    from collections import Counter
    print(f"TOTAL rows: {len(rows)}  by type: {dict(Counter(r['ELEMENT_TYPE'] for r in rows))}")
    print(f"skipped single cells: {stats['skipped_single_cells']}  "
          f"skipped sheets: {stats['skipped_sheets']}  "
          f"sheet errors: {stats['sheet_errors']}")
    print(f"empty-yield sheets: {stats['empty_yield_sheets']}")
    print(f"duplicate scalar names (cross-sheet): {stats['duplicate_scalar_names'][:12]}")
    cur = None
    for r in rows:
        if only and r["SHEET_NAME"] != only:
            continue
        if r["SHEET_NAME"] != cur:
            cur = r["SHEET_NAME"]
            print(f"\n===== {cur} =====")
        rv = r["RAW_VALUE"]
        rv = rv[:32] + "…" if rv and len(rv) > 32 else rv
        print(f"  [{r['ELEMENT_TYPE']:11s}] {(r['ELEMENT_NAME'] or '')[:38]:38s} "
              f"cell={str(r['CELL_REF']):8s} k={str(r['VALUE_KIND'] or ''):8s} "
              f"val={rv!r} {str(r['COLUMNS_JSON'] or '')[:52]}")
