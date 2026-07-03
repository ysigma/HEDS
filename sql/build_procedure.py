#!/usr/bin/env python3
"""Assemble the Snowflake stored procedure from the verified extractor.

The deployable ``extract_template_metadata.sql`` is *generated* — it embeds
``extract_core.py`` verbatim inside a dollar-quoted ``CREATE PROCEDURE`` body
and appends the Snowflake I/O handler. Edit ``extract_core.py`` (the source of
truth) and re-run this script; never hand-edit the ``.sql``.

    python3 sql/build_procedure.py
"""

import os

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "extract_core.py")
OUT = os.path.join(HERE, "extract_template_metadata.sql")

with open(SRC) as f:
    core = f.read()
# Drop the local CLI (`if __name__ == "__main__": ...`); keep everything above.
core = core[: core.index('if __name__ == "__main__":')].rstrip() + "\n"

WRAPPER = r'''

# ==========================================================================
# Snowflake handler: download the template, extract, persist to the metadata
# table. The metadata table schema is self-healing (created if missing, and
# any missing column is added), so this procedure can be deployed over an
# older/naive TEMPLATE_METADATA without a manual migration.
# ==========================================================================

import os
from collections import Counter
from openpyxl import load_workbook
from snowflake.snowpark.types import (
    StructType, StructField, StringType, LongType, TimestampType,
)

TEMPLATES_TABLE = "SE_DEMO_DB.JHI_HEDS_DEMO.REPORT_TEMPLATES"
METADATA_TABLE  = "SE_DEMO_DB.JHI_HEDS_DEMO.TEMPLATE_METADATA"
STAGE_NAME      = "SE_DEMO_DB.JHI_HEDS_DEMO.REPORTS_STAGE"

# (column name, Snowflake type, Snowpark type). Order defines both the DDL and
# the insert DataFrame. Only SHEET_NAME / ELEMENT_TYPE / ELEMENT_NAME /
# CELL_REF / TABLE_REF / COLUMNS_JSON are read by the Sigma plugin (matched by
# name); the rest are extra descriptive metadata it safely ignores.
COLUMN_SPEC = [
    ("TEMPLATE_ID",      "VARCHAR",       StringType()),
    ("ELEMENT_KEY",      "VARCHAR",       StringType()),
    ("ELEMENT_TYPE",     "VARCHAR",       StringType()),
    ("ELEMENT_NAME",     "VARCHAR",       StringType()),
    ("SHEET_NAME",       "VARCHAR",       StringType()),
    ("CELL_REF",         "VARCHAR",       StringType()),
    ("LABEL_REF",        "VARCHAR",       StringType()),
    ("TABLE_REF",        "VARCHAR",       StringType()),
    ("HEADER_REF",       "VARCHAR",       StringType()),
    ("DATA_REF",         "VARCHAR",       StringType()),
    ("ORIENTATION",      "VARCHAR",       StringType()),
    ("NUM_ROWS",         "NUMBER",        LongType()),
    ("NUM_COLS",         "NUMBER",        LongType()),
    ("COLUMNS_JSON",     "VARCHAR",       StringType()),
    ("COL_KINDS_JSON",   "VARCHAR",       StringType()),
    ("VALUE_KIND",       "VARCHAR",       StringType()),
    ("RAW_VALUE",        "VARCHAR",       StringType()),
    ("FORMULA",          "VARCHAR",       StringType()),
    ("DETECTION_METHOD", "VARCHAR",       StringType()),
    ("SEQ",              "NUMBER",        LongType()),
    ("EXTRACTED_AT",     "TIMESTAMP_NTZ", TimestampType()),
]


def ensure_schema(session):
    cols = ", ".join('"%s" %s' % (n, t) for n, t, _ in COLUMN_SPEC)
    session.sql("CREATE TABLE IF NOT EXISTS %s (%s)" % (METADATA_TABLE, cols)).collect()
    for n, t, _ in COLUMN_SPEC:
        try:
            session.sql('ALTER TABLE %s ADD COLUMN IF NOT EXISTS "%s" %s'
                        % (METADATA_TABLE, n, t)).collect()
        except Exception:
            pass


def insert_rows(session, template_id, rows, extracted_at):
    if not rows:
        return 0
    schema = StructType([StructField(n, tp) for n, _, tp in COLUMN_SPEC])
    long_cols = {"NUM_ROWS", "NUM_COLS", "SEQ"}
    data = []
    for r in rows:
        record = []
        for name, _, _ in COLUMN_SPEC:
            if name == "TEMPLATE_ID":
                record.append(template_id)
            elif name == "EXTRACTED_AT":
                record.append(extracted_at)
            elif name in long_cols:
                v = r.get(name)
                record.append(int(v) if v is not None else None)
            else:
                v = r.get(name)
                record.append(None if v is None else str(v))
        data.append(record)
    df = session.create_dataframe(data, schema=schema)
    df.write.mode("append").save_as_table(METADATA_TABLE, column_order="name")
    return len(data)


def main(session, template_id):
    import datetime as _dt
    ensure_schema(session)

    tdf = session.sql(
        "SELECT TEMPLATE_FILE FROM %s WHERE TEMPLATE_ID = ?" % TEMPLATES_TABLE,
        params=[template_id],
    ).to_pandas()
    if tdf.empty:
        return "ERROR: Template ID '%s' not found in REPORT_TEMPLATES" % template_id
    template_file = tdf.iloc[0]["TEMPLATE_FILE"]

    try:
        session.file.get("@%s/templates/%s" % (STAGE_NAME, template_file), "/tmp/")
    except Exception as e:
        return "ERROR: Could not download '%s' from stage: %s" % (template_file, e)

    local_path = os.path.join("/tmp", os.path.basename(template_file))
    try:
        wb = load_workbook(local_path, data_only=False, read_only=False)
        wb_vals = load_workbook(local_path, data_only=True, read_only=False)
    except Exception as e:
        return "ERROR: Could not open workbook '%s': %s" % (template_file, e)

    rows, stats = extract_workbook_metadata(wb, wb_vals)

    extracted_at = _dt.datetime.now()
    session.sql("BEGIN").collect()
    try:
        session.sql("DELETE FROM %s WHERE TEMPLATE_ID = ?" % METADATA_TABLE,
                    params=[template_id]).collect()
        inserted = insert_rows(session, template_id, rows, extracted_at)
        session.sql("COMMIT").collect()
    except Exception as e:
        session.sql("ROLLBACK").collect()
        return "ERROR: Failed to persist metadata for '%s': %s" % (template_id, e)

    summary = {
        "status": "OK",
        "template_id": template_id,
        "template_file": template_file,
        "elements_extracted": inserted,
        "by_type": dict(Counter(r["ELEMENT_TYPE"] for r in rows)),
        "sheets_processed": len({r["SHEET_NAME"] for r in rows}),
        "skipped_single_cells": stats["skipped_single_cells"],
        "skipped_sheets": stats["skipped_sheets"],
        "empty_yield_sheets": stats["empty_yield_sheets"],
        "sheet_errors": stats["sheet_errors"],
        "duplicate_scalar_names": stats["duplicate_scalar_names"],
    }
    return json.dumps(summary)
'''

HEADER = """-- =====================================================================
-- EXTRACT_TEMPLATE_METADATA
--
-- GENERATED FILE - do not edit by hand. Edit sql/extract_core.py and run
-- `python3 sql/build_procedure.py` to regenerate.
--
-- Extracts a rich structural description of an Excel report template into
-- SE_DEMO_DB.JHI_HEDS_DEMO.TEMPLATE_METADATA, for the Sigma report-builder
-- plugin and the downstream population procedure.
--
-- Captures, across every visible worksheet:
--   * Excel Tables (incl. multi-header-row and single-row "wide record" tables)
--   * Defined names (scalars, 1-D lists, 2-D ranges; #REF!/built-ins skipped)
--   * Key/value fact blocks (horizontal & vertical) and inline "Label: value"
--   * Breakdown tables that are NOT Excel Tables (connected-component scan with
--     spacer-column de-fragmentation and section-title peeling)
--   * Footnotes / labeled long text
-- with sheet, cell/table/header/data refs, orientation, row & column counts,
-- header names, number-format-aware value kinds, the cached value and formula.
--
-- The TEMPLATE_METADATA schema is self-healing: created if missing and any
-- absent column is added, so this can be (re)deployed without a migration.
-- The procedure body is dollar-quoted so no single-quote escaping is needed.
-- =====================================================================

CREATE OR REPLACE PROCEDURE SE_DEMO_DB.JHI_HEDS_DEMO.EXTRACT_TEMPLATE_METADATA("TEMPLATE_ID" VARCHAR)
RETURNS VARCHAR
LANGUAGE PYTHON
RUNTIME_VERSION = '3.11'
PACKAGES = ('snowflake-snowpark-python','openpyxl','pandas')
HANDLER = 'main'
EXECUTE AS CALLER
AS $$
"""

FOOTER = "\n$$;\n"


def build():
    body = core + WRAPPER
    assert "$$" not in body, "Python body contains $$ which would break dollar-quoting"
    compile(body, "<procedure>", "exec")  # fail fast if the embedded Python is broken
    with open(OUT, "w") as f:
        f.write(HEADER + body + FOOTER)
    print("Wrote %s (%d bytes); embedded Python compiles OK" % (OUT, len(open(OUT).read())))


if __name__ == "__main__":
    build()
